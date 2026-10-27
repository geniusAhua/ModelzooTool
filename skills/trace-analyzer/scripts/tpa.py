#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
tpa.py — Torch Profile 分析器（单文件，无需解压）

用法:
  python tpa.py analyze  <input>                  # 单组统计（input: trace 文件 / .gz / zip / tar.gz / 目录）
  python tpa.py compare  <ctl_input> <exp_input>  # 对照组 vs 实验组对比（保留原项目功能）
  python tpa.py list     <input>                  # 列出输入里可用的 trace 候选

通用参数:
  --rank N            选择 rank（默认 0；匹配文件名里的 rankN / epN）
  --trace SUBSTR      按文件名子串精确选择 trace
  --decode-num N      decode 采样段数（默认 5）
  --outdir DIR        输出目录（默认 <cwd>/tpa_out）
  --name NAME         产物命名前缀（默认取 trace 文件名主干）
  --top N             Markdown 表格单表行数（默认 15）
  --summary-json PATH  额外输出极简摘要 JSON（给 agent 节约 token 用）

输出: <outdir>/<name>_report.md, <name>.xlsx, (<name>_summary.json)
      compare 模式: <name>_compare.md, <name>_compare.xlsx

依赖: ijson, openpyxl（pandas 不需要）
"""

import argparse
import gzip
import json
import os
import re
import sys
import tarfile
import zipfile
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Dict, Iterable, List, Optional, Tuple

try:
    import ijson
except ImportError:  # pragma: no cover
    sys.exit("缺少依赖 ijson，请先执行: pip install ijson openpyxl")

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError:  # pragma: no cover
    sys.exit("缺少依赖 openpyxl，请先执行: pip install ijson openpyxl")

TRACE_SUFFIXES = (".json", ".json.gz", ".gz", ".pt.trace.json", ".json.gz")
ARCHIVE_SUFFIXES = (".zip", ".tar", ".tar.gz", ".tgz", ".tar.bz2", ".tbz2", ".tar.xz", ".txz")

# ---------------------------------------------------------------------------
# 算子分类规则：顺序敏感，先命中先返回
#   每条: (kernel_class, kernel_type, [小写关键字...])
#   匹配对象 = 规范化后的 kernel 名（已剥离 <...> / (...) 内的模板参数）
# ---------------------------------------------------------------------------
CLASS_RULES: List[Tuple[str, str, Tuple[str, ...]]] = [
    # ---- MoE 相关（含 MoE GEMM，作为一个语义大类）----
    ("moe", "gemm", ("void-mctlass::kernel-moe", "macagemmmoe", "fused_moe_kernel",
                      "fusedmoekernel", "moe_gemm")),
    ("moe", "align_block", ("moe_align_block", "moalignblock")),
    ("moe", "sort_tokens", ("count_and_sort_expert_tokens", "sort_expert")),
    ("moe", "topk", ("single_group_topk", "moe_topk")),
    ("moe", "sum", ("_moe_sum", "moe_sum", "moe_finalize")),
    ("moe", "other", ("moe",)),
    # ---- 通信 ----
    ("communication", "all_reduce", ("allreduce", "cross_device_reduce_1stage",
                                     "cross_device_reduce_2stage", "cross_device_reduce_3stage")),
    ("communication", "all_gather", ("allgather", "all_gather")),
    ("communication", "all_to_all", ("alltoall", "all2all", "all_to_all")),
    ("communication", "send_recv", ("sendrecv", "send_recv", "broadcast")),
    ("communication", "other", ("mcclkernel", "ncclkernel", "mccl", "nccl")),
    # ---- GEMM / GEMV ----
    ("gemm", "gemv", ("gemvt", "gemv")),
    ("gemm", "mctlass", ("void-mctlass::kernel-gemm", "mctlass")),
    ("gemm", "mcblas", ("mcblas",)),
    ("gemm", "cublas", ("cublas", "cutlass")),
    # ---- Attention ----
    ("attention", "kv_cache", ("concat_and_cache", "reshape_and_cache", "gather_and_maybe_dequant",
                               "cache_mla", "kv_cache")),
    ("attention", "attn_merge", ("merge_attn_states", "mergeattnstates")),
    ("attention", "mla", ("_mla_", "mla_kernel", "mla_combine")),
    ("attention", "flash_attn", ("flash_fwd", "flashattn", "flash_attn")),
    ("attention", "paged_attn", ("paged_attention", "pagedattn")),
    # ---- Norm ----
    ("norm", "rms_norm", ("rmsnorm", "rms_norm")),
    ("norm", "layer_norm", ("layernorm", "layer_norm")),
    # ---- Activation ----
    ("activation", "silu_and_mul", ("silu_and_mul", "siluandmul", "act_and_mul", "actandmul")),
    ("activation", "gelu", ("gelu",)),
    ("activation", "other", ("activation", "silu", "sigmoid")),
    # ---- MoE 调度类（非 GEMM 部分：align / sort / topk / sum）----
    # 已在文件头部按语义统一处理，这里不再重复
    # ---- Elementwise / Copy ----
    ("elementwise", "copy", ("direct_copy", "catarraybatchedcopy", "copy_cast")),
    ("elementwise", "add", ("functor_add", "aunaryfunctor", "binaryfunctor")),
    ("elementwise", "fill", ("fillfunctor",)),
    ("elementwise", "other", ("elementwise_kernel", "vectorized_elementwise",
                              "unrolled_elementwise", "elementwise")),
    # ---- Sampling ----
    ("sampling", "topk", ("mbtopk", "bitonicsort", "topk")),
    ("sampling", "other", ("sampling",)),
    # ---- Triton ----
    ("triton", "triton", ("triton",)),
    # ---- Softmax ----
    ("softmax", "softmax", ("softmax",)),
]

STAGE_ORDER = [
    "prefilling",
    "prefilling+decoding",
    "decoding",
    "before-dec",
    "before-pre",
    "after-dec",
    "decoding(excluded)",
]

# 只有连续时间窗的阶段，其“跨度”才有物理意义
CONTIGUOUS_STAGES = {"prefilling", "prefilling+decoding", "decoding"}

RED, GREEN, YELLOW, NC = "\033[0;31m", "\033[0;32m", "\033[33m", "\033[0m"


# ---------------------------------------------------------------------------
# 输入层：文件 / gz / zip / tar(.gz|bz2|xz) / 目录，全部流式打开，不解压到磁盘
# ---------------------------------------------------------------------------
@dataclass
class Candidate:
    label: str          # 展示名（相对路径）
    size: int           # 字节数（压缩后，用于排序）
    opener: object      # () -> binary stream
    rank: Optional[int] = None


def _rank_of(name: str) -> Optional[int]:
    m = re.search(r"rank[_-]?(\d+)", name, re.IGNORECASE)
    if m:
        return int(m.group(1))
    m = re.search(r"\bep(\d+)\b", name, re.IGNORECASE)
    return int(m.group(1)) if m else None


def _is_trace_name(name: str) -> bool:
    low = name.lower()
    if not low.endswith((".json", ".json.gz", ".gz")):
        return False
    return True


def _candidates_of_file(path: Path) -> List[Candidate]:
    name = path.name
    low = name.lower()

    if low.endswith((".zip",)):
        zf = zipfile.ZipFile(path)
        out = []
        for info in zf.infolist():
            if info.is_dir() or not _is_trace_name(info.filename):
                continue
            member = info.filename
            out.append(Candidate(
                label=f"{name}!{member}", size=info.file_size,
                opener=(lambda p=path, m=member: zipfile.ZipFile(p).open(m)),
                rank=_rank_of(member)))
        return out

    if low.endswith(ARCHIVE_SUFFIXES):
        tf = tarfile.open(path, "r:*")
        out = []
        for info in tf.getmembers():
            if not info.isfile() or not _is_trace_name(info.name):
                continue
            member = info.name
            out.append(Candidate(
                label=f"{name}!{member}", size=info.size,
                opener=(lambda p=path, m=member: tarfile.open(p, "r:*").extractfile(m)),
                rank=_rank_of(member)))
        return out

    if low.endswith((".gz",)) and not low.endswith(".tar.gz"):
        return [Candidate(label=name, size=path.stat().st_size,
                          opener=(lambda p=path: gzip.open(p, "rb")), rank=_rank_of(name))]

    return [Candidate(label=name, size=path.stat().st_size,
                      opener=(lambda p=path: open(p, "rb")), rank=_rank_of(name))]


def collect_candidates(target: str) -> List[Candidate]:
    p = Path(target)
    if not p.exists():
        sys.exit(f"{RED}输入不存在: {target}{NC}")
    if p.is_dir():
        out: List[Candidate] = []
        for f in sorted(p.rglob("*")):
            if not f.is_file():
                continue
            low = f.name.lower()
            if low.endswith(ARCHIVE_SUFFIXES + (".gz", ".json")):
                try:
                    out.extend(_candidates_of_file(f))
                except Exception as e:  # 损坏的归档不应中断整体扫描
                    print(f"{YELLOW}跳过无法读取的归档 {f}: {e}{NC}")
        return out
    return _candidates_of_file(p)


def _pick_default(pool: List[Candidate]) -> Candidate:
    """默认候选：排除 async_llm → 优先 rank0 → 否则 rank 最小 → 否则文件最大"""
    gpu = [c for c in pool if "async_llm" not in c.label.lower()]
    pool = gpu or pool
    rank0 = [c for c in pool if c.rank == 0]
    if rank0:
        return max(rank0, key=lambda c: c.size)
    ranked = [c for c in pool if c.rank is not None]
    if ranked:
        return min(ranked, key=lambda c: (c.rank, -c.size))
    return max(pool, key=lambda c: c.size)


def select_candidate(cands: List[Candidate], rank: Optional[int], trace_substr: Optional[str]) -> Candidate:
    if not cands:
        sys.exit(f"{RED}未找到 traceEvents 文件（支持 .json/.json.gz/zip/tar.gz/目录）{NC}")
    pool = cands
    if trace_substr:
        pool = [c for c in pool if trace_substr in c.label]
        if not pool:
            sys.exit(f"{RED}--trace {trace_substr} 未匹配到任何文件{NC}")

    if rank is not None:
        exact = [c for c in pool if c.rank == rank]
        if exact:
            return max(exact, key=lambda c: c.size)
        print(f"{YELLOW}未找到 rank={rank} 的 trace，回退到默认候选{NC}")

    return _pick_default(pool)


def cmd_list(target: str, as_json: bool = False) -> None:
    cands = collect_candidates(target)
    default = _pick_default(cands)
    rows = []
    for c in sorted(cands, key=lambda x: (x.rank is None,
                                          x.rank if x.rank is not None else 0,
                                          -x.size)):
        rows.append({
            "name": c.label,
            "rank": c.rank,
            "size_mb": round(c.size / 1024 / 1024, 2),
            "is_async_llm": "async_llm" in c.label.lower(),
            "recommended": c is default,
        })

    if as_json:
        print(json.dumps({
            "input": target,
            "count": len(rows),
            "gpu_rank_count": sum(1 for r in rows if not r["is_async_llm"]),
            "recommended": {"name": default.label, "rank": default.rank},
            "candidates": rows,
        }, ensure_ascii=False, indent=2))
        return

    print(f"共 {len(cands)} 个候选 trace（推荐 {('rank' + str(default.rank)) if default.rank is not None else default.label}）：")
    for r in rows:
        mark = " *" if r["recommended"] else "  "
        note = "  <= 推荐" if r["recommended"] else ("  <= async_llm，默认排除" if r["is_async_llm"] else "")
        print(f"{mark} rank={str(r['rank']):>4}  {r['size_mb']:8.2f}MB  {r['name']}{note}")
    print(f"\n默认使用: {default.label}")
    print("提示：多 rank 场景请先询问用户（默认推荐 rank0），再传 --rank N 固定选择。")


# ---------------------------------------------------------------------------
# 解析层
# ---------------------------------------------------------------------------
@dataclass
class Kernel:
    raw: str
    norm: str
    ts: float
    dur: float


@dataclass
class TraceData:
    source: str
    kernels: List[Kernel] = field(default_factory=list)
    annotations: List[Tuple[str, float, float]] = field(default_factory=list)
    event_count: int = 0
    cat_counter: Counter = field(default_factory=Counter)
    streams: set = field(default_factory=set)


class _Closer:
    """统一的流关闭器：主数据流 + 需要连带关闭的底层容器"""

    def __init__(self, stream, extras: Iterable = ()):
        self.stream = stream
        self.extras = list(extras)

    def close(self) -> None:
        try:
            self.stream.close()
        finally:
            for c in self.extras:
                try:
                    c.close()
                except Exception:
                    pass


def _open_stream(cand: Candidate) -> _Closer:
    """打开候选 trace 数据流。

    - 本地 .gz 文件：opener 已是 gzip.open，不需要额外处理
    - zip / tar 成员：如果成员自身还是 .gz（torch 默认导出格式），这里再包一层 gzip
    """
    stream = cand.opener()
    extras: List = []
    is_member = "!" in cand.label
    inner = cand.label.rsplit("!", 1)[-1].lower()
    if is_member and inner.endswith(".gz"):
        gz = gzip.GzipFile(fileobj=stream)
        extras.append(stream)
        stream = gz
    return _Closer(stream, extras)


def normalize_kernel_name(raw: str) -> str:
    """去掉 <...> 与 (...) 内的模板参数/实参，只保留可读的主干名"""
    if not raw:
        return ""
    out, depth = [], 0
    for ch in raw:
        if ch in "<(":
            depth += 1
        elif ch in ">)":
            if depth > 0:
                depth -= 1
        elif depth == 0:
            out.append(ch)
    norm = "".join(out).strip().replace(" ", "-")
    if raw.startswith("void mctlass::Kernel"):
        if "MacaGemmMoe" in raw:
            norm += "-moe"
        else:
            norm += "-gemm"
    return norm


def parse_trace(cand: Candidate) -> TraceData:
    td = TraceData(source=cand.label)

    def run(prefix: str) -> None:
        closer = _open_stream(cand)
        try:
            for ev in ijson.items(closer.stream, prefix, use_float=True):
                td.event_count += 1
                if not isinstance(ev, dict):
                    continue
                cat = ev.get("cat")
                td.cat_counter[cat] += 1
                if cat not in ("kernel", "gpu_user_annotation"):
                    continue
                name = ev.get("name")
                ts, dur = ev.get("ts"), ev.get("dur")
                if name is None or ts is None or dur is None:
                    continue
                if cat == "kernel":
                    td.kernels.append(Kernel(name, normalize_kernel_name(name), ts, dur))
                    td.streams.add((ev.get("pid"), ev.get("tid")))
                elif name.startswith("execute_"):
                    td.annotations.append((name, ts, dur))
        finally:
            closer.close()

    try:
        run("traceEvents.item")
    except ijson.common.JSONError:
        # 兜底：顶层直接是事件数组的 trace
        td.kernels.clear()
        td.annotations.clear()
        td.event_count = 0
        td.cat_counter.clear()
        td.streams.clear()
        run("item")

    td.kernels.sort(key=lambda k: k.ts)
    td.annotations.sort(key=lambda a: a[1])
    return td


# ---------------------------------------------------------------------------
# 阶段划分
# ---------------------------------------------------------------------------
@dataclass
class Segment:
    kind: str      # p / d / pd
    start: float
    end: float
    index: int = 0


def _ann_kind(name: str) -> Optional[str]:
    m1 = re.search(r"context[_(](\d+)\s*[()]", name)
    m2 = re.search(r"generation[_(](\d+)\s*[()]", name)
    if m1 and m2:
        ctx, gen = int(m1.group(1)), int(m2.group(1))
    else:  # 兼容 vllm execute_context_1(12288)_generation_0(0) 的 split 写法
        parts = name.split("_")
        try:
            ctx = int(parts[2].split("(")[0])
            gen = int(parts[4].split("(")[0])
        except Exception:
            return None
    if ctx > 0 and gen == 0:
        return "p"
    if ctx == 0 and gen > 0:
        return "d"
    if ctx > 0 and gen > 0:
        return "pd"
    return None


def build_segments(annotations: List[Tuple[str, float, float]]) -> Tuple[List[Segment], List[str]]:
    warnings: List[str] = []
    segs: List[Segment] = []
    for name, ts, dur in annotations:
        kind = _ann_kind(name)
        if kind is None or dur <= 0:
            continue
        start, end = ts, ts + dur
        if segs and start < segs[-1].end:
            prev = segs[-1]
            if prev.kind == kind:
                prev.end = max(prev.end, end)   # 同名同阶段重叠 → 合并
                continue
            # 不同阶段重叠：裁剪，避免 kernel 双重归属
            warnings.append(f"注解重叠且阶段不同（{prev.kind} vs {kind}），已裁剪至不重叠")
            start = prev.end
            if start >= end:
                continue
        segs.append(Segment(kind=kind, start=start, end=end))
    for i, s in enumerate(segs):
        s.index = i
    return segs, warnings


def assign_stages(td: TraceData, decode_num: int) -> Tuple[List[Tuple[Kernel, str]], dict]:
    """给每个 kernel 打 stage 标签；返回 [(kernel, stage)] 与阶段元信息"""
    segs, warnings = build_segments(td.annotations)
    if not segs:
        raise SystemExit(f"{RED}未找到可识别的 execute_context_*_generation_* 注解，"
                         f"无法划分 prefill/decode 阶段{NC}")

    d_segs = [s for s in segs if s.kind == "d"]
    if len(d_segs) < decode_num:
        raise SystemExit(
            f"{RED}decode 阶段只有 {len(d_segs)} 个，少于需要的 {decode_num} 个。{NC}\n"
            f"提示：profile 抓取步数太少，请把 profile 的 output/steps 设大（推荐 32）后重抓；"
            f"或用 --decode-num {len(d_segs)} 降低采样数。")
    sampled_d = set(s.index for s in d_segs[-decode_num:])

    kind_label = {"p": "prefilling", "pd": "prefilling+decoding", "d": "decoding"}
    for s in segs:
        if s.kind == "d" and s.index not in sampled_d:
            s.kind = "d_old"

    active = [s for s in segs if s.kind != "d_old"]
    win_start = active[0].start
    win_end = active[-1].end

    def label_of(seg: Segment) -> str:
        if seg.kind == "d_old":
            return "decoding(excluded)"
        return kind_label[seg.kind]

    staged: List[Tuple[Kernel, str]] = []
    si = 0
    for k in td.kernels:
        # 指针推进（kernel 已按 ts 排序）
        while si < len(segs) and k.ts >= segs[si].end:
            si += 1
        stage = None
        if si < len(segs) and k.ts >= segs[si].start:
            stage = label_of(segs[si])
        elif k.ts < win_start:
            stage = "before-pre"
        elif k.ts >= win_end:
            stage = "after-dec"
        else:
            stage = "before-dec"
        staged.append((k, stage))

    meta = {
        "segments": segs,
        "prefill_segments": [s for s in segs if s.kind == "p"],
        "pd_segments": [s for s in segs if s.kind == "pd"],
        "decode_segments": [s for s in segs if s.kind == "d"],
        "decode_excluded": [s for s in segs if s.kind == "d_old"],
        "decode_num": decode_num,
        "warnings": warnings,
    }
    return staged, meta


# ---------------------------------------------------------------------------
# 分类 + 汇总
# ---------------------------------------------------------------------------
def classify_kernel(k: Kernel) -> Tuple[str, str]:
    low = k.norm.lower()
    for cls, typ, kws in CLASS_RULES:
        for kw in kws:
            if kw in low:
                return cls, typ
    return "other", "other"


def _pct(num: float, den: float) -> float:
    return (num / den * 100.0) if den else 0.0


def _class_split_us(a: "Analysis", cls: str) -> Tuple[float, float, float]:
    """某个大类在 (prefilling, decoding 采样段, prefilling+decoding 混合段) 的耗时"""
    return (a.stage_class_us.get(("prefilling", cls), 0.0),
            a.stage_class_us.get(("decoding", cls), 0.0),
            a.stage_class_us.get(("prefilling+decoding", cls), 0.0))


@dataclass
class Analysis:
    name: str
    source: str
    trace: TraceData
    staged: List[Tuple[Kernel, str]]
    meta: dict
    gpu_us: float = 0.0
    span_us: float = 0.0
    kernel_count: int = 0
    stage_us: Dict[str, float] = field(default_factory=dict)
    stage_cnt: Dict[str, int] = field(default_factory=dict)
    stage_span: Dict[str, float] = field(default_factory=dict)
    class_us: Dict[str, float] = field(default_factory=dict)
    class_cnt: Dict[str, int] = field(default_factory=dict)
    type_us: Dict[Tuple[str, str], float] = field(default_factory=dict)
    type_cnt: Dict[Tuple[str, str], int] = field(default_factory=dict)
    kernel_us: Dict[Tuple[str, str, str], float] = field(default_factory=dict)
    kernel_cnt: Dict[Tuple[str, str, str], int] = field(default_factory=dict)
    stage_class_us: Dict[Tuple[str, str], float] = field(default_factory=dict)
    stage_class_cnt: Dict[Tuple[str, str], int] = field(default_factory=dict)
    stage_type_us: Dict[Tuple[str, str, str], float] = field(default_factory=dict)
    stage_type_cnt: Dict[Tuple[str, str, str], int] = field(default_factory=dict)
    stage_kernel_us: Dict[Tuple[str, str, str, str], float] = field(default_factory=dict)
    stage_kernel_cnt: Dict[Tuple[str, str, str, str], int] = field(default_factory=dict)


def analyze(cand: Candidate, name: str, decode_num: int) -> Analysis:
    print(f"解析 {cand.label} ...", flush=True)
    td = parse_trace(cand)
    print(f"  事件 {td.event_count}，kernel {len(td.kernels)}，"
          f"execute 注解 {len(td.annotations)}，stream {len(td.streams)}", flush=True)
    staged, meta = assign_stages(td, decode_num)

    a = Analysis(name=name, source=cand.label, trace=td, staged=staged, meta=meta)
    a.kernel_count = len(staged)
    a.gpu_us = sum(k.dur for k, _ in staged)
    if staged:
        first = min(k.ts for k, _ in staged)
        last = max(k.ts + k.dur for k, _ in staged)
        a.span_us = last - first

    span_acc: Dict[str, List[float]] = defaultdict(list)
    for k, stage in staged:
        cls, typ = classify_kernel(k)
        a.stage_us[stage] = a.stage_us.get(stage, 0.0) + k.dur
        a.stage_cnt[stage] = a.stage_cnt.get(stage, 0) + 1
        a.class_us[cls] = a.class_us.get(cls, 0.0) + k.dur
        a.class_cnt[cls] = a.class_cnt.get(cls, 0) + 1
        a.type_us[(cls, typ)] = a.type_us.get((cls, typ), 0.0) + k.dur
        a.type_cnt[(cls, typ)] = a.type_cnt.get((cls, typ), 0) + 1
        key = (cls, typ, k.norm)
        a.kernel_us[key] = a.kernel_us.get(key, 0.0) + k.dur
        a.kernel_cnt[key] = a.kernel_cnt.get(key, 0) + 1
        a.stage_class_us[(stage, cls)] = a.stage_class_us.get((stage, cls), 0.0) + k.dur
        a.stage_class_cnt[(stage, cls)] = a.stage_class_cnt.get((stage, cls), 0) + 1
        a.stage_type_us[(stage, cls, typ)] = a.stage_type_us.get((stage, cls, typ), 0.0) + k.dur
        a.stage_type_cnt[(stage, cls, typ)] = a.stage_type_cnt.get((stage, cls, typ), 0) + 1
        sk = (stage, cls, typ, k.norm)
        a.stage_kernel_us[sk] = a.stage_kernel_us.get(sk, 0.0) + k.dur
        a.stage_kernel_cnt[sk] = a.stage_kernel_cnt.get(sk, 0) + 1
        span_acc[stage].append(k.ts)
        span_acc[stage].append(k.ts + k.dur)

    for stage, ts in span_acc.items():
        a.stage_span[stage] = max(ts) - min(ts)
    return a


def ordered_stages(a: Analysis) -> List[str]:
    known = [s for s in STAGE_ORDER if s in a.stage_us]
    extra = sorted(s for s in a.stage_us if s not in STAGE_ORDER)
    return known + extra


def top_items(d: Dict, n: int) -> List[Tuple]:
    return sorted(d.items(), key=lambda kv: -kv[1])[:n]


# ---------------------------------------------------------------------------
# 输出：Excel
# ---------------------------------------------------------------------------
HDR_FILL = PatternFill("solid", fgColor="DDEBF7")
GRP_FILL = PatternFill("solid", fgColor="F2F2F2")
TOTAL_FILL = PatternFill("solid", fgColor="FFF2CC")


def _style_header(ws, widths: Optional[List[int]] = None) -> None:
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = HDR_FILL
        c.alignment = Alignment(horizontal="center", vertical="center")
    ws.freeze_panes = "A2"
    if widths:
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w


def _sheet_summary(wb, a: Analysis) -> None:
    ws = wb.create_sheet("Summary")
    ws.append(["项目", "值", "说明"])
    rows = [
        ("trace 来源", a.source, "所选 trace 文件（压缩包内成员以 ! 分隔）"),
        ("kernel 事件数", a.kernel_count, "cat=kernel 的事件总数"),
        ("trace 事件总数", a.trace.event_count, "全部 traceEvents 条目"),
        ("GPU 耗时合计 (ms)", round(a.gpu_us / 1000, 3), "所有 kernel dur 累加"),
        ("profile 覆盖时长 (ms)", round(a.span_us / 1000, 3), "首个 kernel start → 末个 kernel end"),
        ("GPU 忙碌率 (%)", round(_pct(a.gpu_us, a.span_us), 2), "GPU 耗时合计 / 覆盖时长"),
        ("kernel 并发流数", len(a.trace.streams), ">1 时累加耗时会含并发重叠"),
        ("prefill 段数", len(a.meta["prefill_segments"]), "chunked prefill 会有多段，全部计入 prefilling"),
        ("prefill+decode 混合段数", len(a.meta["pd_segments"]), "ctx>0 且 gen>0 的段"),
        ("decode 采样段数", a.meta["decode_num"], "只统计最后 N 个纯 decode 段（默认 5）"),
        ("decode 未采样段数", len(a.meta["decode_excluded"]), "更早的纯 decode 段，仅计入全量统计"),
        ("生成时间", datetime.now().strftime("%Y-%m-%d %H:%M:%S"), ""),
    ]
    for r in rows:
        ws.append(list(r))
    _style_header(ws, [26, 62, 46])


def _sheet_stages(wb, a: Analysis) -> None:
    ws = wb.create_sheet("Stages")
    ws.append(["stage", "kernel_count", "kernel_us", "kernel_ms", "span_us",
               "pct_of_gpu", "pct_of_span", "note"])
    note_map = {
        "prefilling": "全部纯 prefill 段（chunked prefill 合并统计）",
        "prefilling+decoding": "ctx>0 且 gen>0 的混合段",
        "decoding": f"仅最后 {a.meta['decode_num']} 个纯 decode 段（采样）",
        "before-dec": "decode 采样窗口内、不属于任何阶段的 kernel",
        "before-pre": "首个阶段之前的 kernel",
        "after-dec": "最后一个采样阶段之后的 kernel",
        "decoding(excluded)": "更早的纯 decode 段，不进入阶段口径，仅计入全量",
    }
    for s in ordered_stages(a):
        us = a.stage_us[s]
        ws.append([s, a.stage_cnt[s], round(us, 3), round(us / 1000, 3),
                   round(a.stage_span.get(s, 0.0), 3) if s in CONTIGUOUS_STAGES else "",
                   round(_pct(us, a.gpu_us), 2), round(_pct(us, a.span_us), 2),
                   note_map.get(s, "")])
    ws.append(["TOTAL", sum(a.stage_cnt.values()), round(a.gpu_us, 3),
               round(a.gpu_us / 1000, 3), round(a.span_us, 3), 100.0,
               round(_pct(a.gpu_us, a.span_us), 2), "各阶段之和 == GPU 耗时合计"])
    for c in ws[ws.max_row]:
        c.fill = TOTAL_FILL
    _style_header(ws, [22, 14, 14, 14, 14, 12, 12, 60])


def _sheet_classes(wb, a: Analysis) -> None:
    ws = wb.create_sheet("Classes")
    ws.append(["kernel_class", "kernel_count", "time_us", "time_ms", "pct_of_gpu",
               "prefill_us", "prefill_ms", "pct_of_prefill",
               "decode_us", "decode_ms", "pct_of_decode",
               "pd_us", "pd_ms", "main_stage"])
    pre_total = a.stage_us.get("prefilling", 0.0)
    dec_total = a.stage_us.get("decoding", 0.0)
    for cls, us in sorted(a.class_us.items(), key=lambda kv: -kv[1]):
        pre, dec, pd = _class_split_us(a, cls)
        by_stage = {s: a.stage_class_us.get((s, cls), 0.0) for s in a.stage_us}
        main_stage = max(by_stage.items(), key=lambda kv: kv[1])[0] if by_stage else ""
        ws.append([cls, a.class_cnt[cls], round(us, 3), round(us / 1000, 3),
                   round(_pct(us, a.gpu_us), 2),
                   round(pre, 3), round(pre / 1000, 3), round(_pct(pre, pre_total), 2),
                   round(dec, 3), round(dec / 1000, 3), round(_pct(dec, dec_total), 2),
                   round(pd, 3), round(pd / 1000, 3), main_stage])
    pre_all = a.stage_us.get("prefilling", 0.0)
    dec_all = a.stage_us.get("decoding", 0.0)
    pd_all = a.stage_us.get("prefilling+decoding", 0.0)
    ws.append(["TOTAL", sum(a.class_cnt.values()), round(a.gpu_us, 3),
               round(a.gpu_us / 1000, 3), 100.0,
               round(pre_all, 3), round(pre_all / 1000, 3), 100.0,
               round(dec_all, 3), round(dec_all / 1000, 3), 100.0,
               round(pd_all, 3), round(pd_all / 1000, 3), ""])
    for c in ws[ws.max_row]:
        c.fill = TOTAL_FILL
    _style_header(ws, [28, 14, 14, 14, 12, 14, 14, 14, 14, 14, 14, 14, 14, 22])


def _sheet_types(wb, a: Analysis) -> None:
    ws = wb.create_sheet("Types")
    ws.append(["kernel_class", "kernel_type", "kernel_count", "time_us", "time_ms",
               "pct_of_class", "pct_of_gpu", "main_stage"])
    for (cls, typ), us in sorted(a.type_us.items(), key=lambda kv: (kv[0][0], -kv[1])):
        by_stage = {s: a.stage_type_us.get((s, cls, typ), 0.0) for s in a.stage_us}
        main_stage = max(by_stage.items(), key=lambda kv: kv[1])[0] if by_stage else ""
        ws.append([cls, typ, a.type_cnt[(cls, typ)], round(us, 3), round(us / 1000, 3),
                   round(_pct(us, a.class_us[cls]), 2), round(_pct(us, a.gpu_us), 2),
                   main_stage])
    ws.append(["TOTAL", "", sum(a.type_cnt.values()), round(a.gpu_us, 3),
               round(a.gpu_us / 1000, 3), "", 100.0, ""])
    for c in ws[ws.max_row]:
        c.fill = TOTAL_FILL
    _style_header(ws, [22, 22, 14, 14, 14, 14, 12, 20])


def _sheet_kernels(wb, a: Analysis) -> None:
    ws = wb.create_sheet("Kernels")
    ws.append(["kernel_class", "kernel_type", "kernel_name", "kernel_count", "time_us",
               "time_ms", "pct_of_gpu", "main_stage", "stages"])
    per_kernel_stage: Dict[Tuple[str, str, str], Counter] = defaultdict(Counter)
    for k, stage in a.staged:
        cls, typ = classify_kernel(k)
        per_kernel_stage[(cls, typ, k.norm)][stage] += k.dur
    for key, us in sorted(a.kernel_us.items(), key=lambda kv: -kv[1]):
        cls, typ, nm = key
        stages = per_kernel_stage[key]
        main_stage = stages.most_common(1)[0][0] if stages else ""
        ws.append([cls, typ, nm, a.kernel_cnt[key], round(us, 3), round(us / 1000, 3),
                   round(_pct(us, a.gpu_us), 2), main_stage,
                   ", ".join(f"{s}:{round(v/1000,1)}ms" for s, v in stages.most_common())])
    _style_header(ws, [22, 22, 70, 14, 14, 14, 12, 20, 60])


def _sheet_hierarchy(wb, a: Analysis) -> None:
    """原项目风格的层级表：stage → class → type → kernel，父行合计"""
    ws = wb.create_sheet("Hierarchy")
    ws.append(["level", "stage", "kernel_class", "kernel_type", "kernel_name",
               "kernel_count", "time_us", "time_ms", "pct_of_stage",
               "pct_of_class", "pct_of_type", "pct_of_gpu"])

    def emit(level, stage, cls, typ, nm, cnt, us, p_stage, p_class=None, p_type=None):
        ws.append([level, stage, cls, typ, nm, cnt, round(us, 3), round(us / 1000, 3),
                   round(p_stage, 4) if p_stage is not None else "",
                   round(p_class, 4) if p_class is not None else "",
                   round(p_type, 4) if p_type is not None else "",
                   round(_pct(us, a.gpu_us), 4)])
        if level == 0:
            for c in ws[ws.max_row]:
                c.fill = GRP_FILL
                c.font = Font(bold=True)
        elif nm == "total":
            for c in ws[ws.max_row]:
                c.fill = TOTAL_FILL

    for stage in ordered_stages(a):
        s_us = a.stage_us[stage]
        emit(0, stage, "total", "total", "total", a.stage_cnt[stage], s_us, 100.0)
        for (cls, us) in sorted(((c, v) for (s, c), v in a.stage_class_us.items() if s == stage),
                                key=lambda kv: -kv[1]):
            emit(1, stage, cls, "total", "total",
                 a.stage_class_cnt[(stage, cls)], us, _pct(us, s_us), 100.0)
            for (c2, t), t_us in sorted(((k, v) for k, v in a.type_us.items() if k[0] == cls),
                                        key=lambda kv: -kv[1]):
                st_us = a.stage_type_us.get((stage, c2, t), 0.0)
                if st_us <= 0:
                    continue
                st_cnt = a.stage_type_cnt.get((stage, c2, t), 0)
                emit(2, stage, cls, t, "total", st_cnt, st_us,
                     _pct(st_us, s_us), _pct(st_us, us), 100.0)
                for (c3, t3, nm), k_us in sorted(
                        ((k, v) for k, v in a.kernel_us.items()
                         if k[0] == cls and k[1] == t),
                        key=lambda kv: -kv[1]):
                    k_stage_us = a.stage_kernel_us.get((stage, c3, t3, nm), 0.0)
                    if k_stage_us <= 0:
                        continue
                    k_stage_cnt = a.stage_kernel_cnt.get((stage, c3, t3, nm), 0)
                    emit(3, stage, cls, t, nm, k_stage_cnt, k_stage_us,
                         _pct(k_stage_us, s_us), _pct(k_stage_us, us),
                         _pct(k_stage_us, t_us))
    _style_header(ws, [8, 22, 20, 20, 70, 14, 14, 14, 12, 12, 12, 12])


def write_excel(a: Analysis, out: Path) -> None:
    wb = Workbook()
    wb.remove(wb.active)
    _sheet_summary(wb, a)
    _sheet_stages(wb, a)
    _sheet_classes(wb, a)
    _sheet_types(wb, a)
    _sheet_kernels(wb, a)
    _sheet_hierarchy(wb, a)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)


# ---------------------------------------------------------------------------
# 输出：Markdown
# ---------------------------------------------------------------------------
def _fmt_ms(us: float) -> str:
    return f"{us/1000:,.2f}"


def write_markdown(a: Analysis, out: Path, top: int) -> None:
    n = a.meta["decode_num"]
    dec_us = a.stage_us.get("decoding", 0.0)
    dec_segs = a.meta["decode_segments"]
    pre_segs = a.meta["prefill_segments"]
    dec_avg = dec_us / n if n else 0.0
    dec_step_ms = (sum(s.end - s.start for s in dec_segs) / n / 1000) if n else 0.0

    L: List[str] = []
    L.append(f"# Torch Profile 分析报告：{a.name}")
    L.append("")
    L.append(f"> 数据来源：`{a.source}`  ")
    L.append(f"> 生成时间：{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}  ")
    L.append(f"> 分析工具：`tpa.py`（ModelzooTool `trace-analyzer` skill）")
    L.append("")

    L.append("## 1. 整体 GPU 耗时")
    L.append("")
    L.append("| 指标 | 数值 |")
    L.append("| --- | ---: |")
    L.append(f"| GPU 耗时合计（所有 kernel 累加） | **{_fmt_ms(a.gpu_us)} ms** |")
    L.append(f"| profile 覆盖时长 | {_fmt_ms(a.span_us)} ms |")
    L.append(f"| GPU 忙碌率 | {_pct(a.gpu_us, a.span_us):.2f}% |")
    L.append(f"| kernel 事件数 | {a.kernel_count:,} |")
    L.append(f"| trace 事件总数 | {a.trace.event_count:,} |")
    L.append(f"| kernel 并发流数 | {len(a.trace.streams)} |")
    L.append("")
    if len(a.trace.streams) > 1:
        L.append("> ⚠️ 检测到多个并发流，kernel 耗时累加会包含并发重叠，"
                 "不能直接当作 wall-clock 时间；wall-clock 请参考“profile 覆盖时长”。")
        L.append("")

    L.append("## 2. 各阶段耗时")
    L.append("")
    L.append(f"- prefill：共 **{len(pre_segs)} 段**（chunked prefill 全部计入 `prefilling`）"
             f"，合计 **{_fmt_ms(a.stage_us.get('prefilling', 0.0))} ms**")
    if a.meta["pd_segments"]:
        L.append(f"- prefill+decode 混合段：**{len(a.meta['pd_segments'])} 段**，"
                 f"合计 {_fmt_ms(a.stage_us.get('prefilling+decoding', 0.0))} ms")
    L.append(f"- decode：**采样最后 {n} 个纯 decode 段**，合计 **{_fmt_ms(dec_us)} ms**，"
             f"平均 {_fmt_ms(dec_avg)} ms/段（单步 wall-clock ≈ {dec_step_ms:.2f} ms）")
    if a.meta["decode_excluded"]:
        L.append(f"- 另有 {len(a.meta['decode_excluded'])} 个更早的 decode 段未纳入阶段统计"
                 f"（仅计入全量 GPU 耗时）")
    L.append("")
    L.append("| stage | kernel 数 | 耗时 (ms) | 占 GPU 总耗时 | 阶段跨度 (ms) | 说明 |")
    L.append("| --- | ---: | ---: | ---: | ---: | --- |")
    note_map = {
        "prefilling": f"全部纯 prefill（{len(pre_segs)} 段）",
        "prefilling+decoding": "ctx>0 且 gen>0 的混合段",
        "decoding": f"最后 {n} 个纯 decode（采样）",
        "before-dec": "采样窗口内非阶段 kernel",
        "before-pre": "首个阶段之前的 kernel",
        "after-dec": "最后采样阶段之后的 kernel",
        "decoding(excluded)": "更早 decode 段（不计入阶段口径）",
    }
    for s in ordered_stages(a):
        span_txt = _fmt_ms(a.stage_span.get(s, 0.0)) if s in CONTIGUOUS_STAGES else "-"
        L.append(f"| {s} | {a.stage_cnt[s]:,} | {_fmt_ms(a.stage_us[s])} | "
                 f"{_pct(a.stage_us[s], a.gpu_us):.2f}% | "
                 f"{span_txt} | {note_map.get(s, '')} |")
    L.append(f"| **合计** | **{sum(a.stage_cnt.values()):,}** | **{_fmt_ms(a.gpu_us)}** | **100%** | "
             f"{_fmt_ms(a.span_us)} | 各行之和 == GPU 耗时合计 |")
    L.append("")

    L.append("## 3. 算子大类耗时（prefill / decode 拆分）")
    L.append("")
    L.append(f"- prefill 列 = `prefilling` 阶段（{len(pre_segs)} 段）内该大类的耗时")
    L.append(f"- decode 列 = `decoding` 阶段（最后 {n} 段采样）内该大类的耗时")
    L.append("")
    L.append("| 大类 (kernel_class) | 总耗时 (ms) | 占 GPU | prefill (ms) | 占 prefill | "
             "decode (ms) | 占 decode |")
    L.append("| --- | ---: | ---: | ---: | ---: | ---: | ---: |")
    pre_total = a.stage_us.get("prefilling", 0.0)
    dec_total = a.stage_us.get("decoding", 0.0)
    for cls, us in top_items(a.class_us, top):
        pre, dec, _ = _class_split_us(a, cls)
        L.append(f"| {cls} | {_fmt_ms(us)} | {_pct(us, a.gpu_us):.2f}% | "
                 f"{_fmt_ms(pre)} | {_pct(pre, pre_total):.2f}% | "
                 f"{_fmt_ms(dec)} | {_pct(dec, dec_total):.2f}% |")
    others = sorted(a.class_us.items(), key=lambda kv: -kv[1])[top:]
    if others:
        rest = sum(v for _, v in others)
        rest_pre = sum(_class_split_us(a, c)[0] for c, _ in others)
        rest_dec = sum(_class_split_us(a, c)[1] for c, _ in others)
        L.append(f"| （其余 {len(others)} 类） | {_fmt_ms(rest)} | {_pct(rest, a.gpu_us):.2f}% | "
                 f"{_fmt_ms(rest_pre)} | {_pct(rest_pre, pre_total):.2f}% | "
                 f"{_fmt_ms(rest_dec)} | {_pct(rest_dec, dec_total):.2f}% |")
    pd_us = a.stage_us.get("prefilling+decoding", 0.0)
    if pd_us:
        L.append("")
        L.append(f"> 另有 **{_fmt_ms(pd_us)} ms**（占 GPU {_pct(pd_us, a.gpu_us):.2f}%）落在 "
                 f"`prefilling+decoding` 混合段，未计入上面 prefill / decode 两列；"
                 f"各类在混合段内的明细见 Excel `Classes` 的 `pd_ms` 列。")
    L.append("")

    L.append("## 4. 算子子类耗时（Top %d）" % top)
    L.append("")
    L.append("| 大类 | 子类 (kernel_type) | 耗时 (ms) | 占本类 | 占 GPU |")
    L.append("| --- | --- | ---: | ---: | ---: |")
    for (cls, typ), us in top_items(a.type_us, top):
        L.append(f"| {cls} | {typ} | {_fmt_ms(us)} | {_pct(us, a.class_us[cls]):.2f}% | "
                 f"{_pct(us, a.gpu_us):.2f}% |")
    L.append("")

    L.append("## 5. 单算子耗时（Top %d）" % top)
    L.append("")
    L.append("| 大类 / 子类 | 算子 | 耗时 (ms) | 占 GPU | 次数 |")
    L.append("| --- | --- | ---: | ---: | ---: |")
    for (cls, typ, nm), us in top_items(a.kernel_us, top):
        short = nm if len(nm) <= 64 else nm[:61] + "..."
        L.append(f"| {cls} / {typ} | `{short}` | {_fmt_ms(us)} | {_pct(us, a.gpu_us):.2f}% | "
                 f"{a.kernel_cnt[(cls, typ, nm)]:,} |")
    L.append("")

    L.append("## 6. 统计口径与校验")
    L.append("")
    L.append("1. kernel 来源：trace 中 `cat == \"kernel\"` 的事件，耗时字段为 `dur`（微秒）。")
    L.append("2. 阶段来源：`cat == \"gpu_user_annotation\"` 且 name 以 `execute_` 开头的注解，"
             "按 `context_N(...)` / `generation_M(...)` 判定：`N>0 且 M==0` → prefill，"
             "`N==0 且 M>0` → decode，两者皆 >0 → 混合段。")
    L.append(f"3. **decode 采样数量为 {n}**：只取时间线上最后 {n} 个纯 decode 段；"
             f"更早的 decode 段标记为 `decoding(excluded)`，只计入全量统计。")
    L.append(f"4. prefill 统计全部 {len(pre_segs)} 段纯 prefill 注解（chunked prefill 会拆成多段）。")
    L.append("5. 落在注解时间窗之外的 kernel 归入 `before-pre` / `before-dec` / `after-dec`，"
             "不会丢失。")
    L.append("6. 校验：各 stage 的 kernel 耗时之和 = GPU 耗时合计 = "
             f"{_fmt_ms(a.gpu_us)} ms（逐 kernel 唯一归属，不重不漏）。")
    if a.meta["warnings"]:
        L.append("")
        L.append("**解析告警：**")
        for w in sorted(set(a.meta["warnings"])):
            L.append(f"- {w}")
    L.append("")

    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text("\n".join(L), encoding="utf-8")


def _summary_payload(a: Analysis, top: int = 5) -> dict:
    return {
        "name": a.name,
        "source": a.source,
        "gpu_kernel_ms": round(a.gpu_us / 1000, 3),
        "span_ms": round(a.span_us / 1000, 3),
        "busy_pct": round(_pct(a.gpu_us, a.span_us), 2),
        "kernel_count": a.kernel_count,
        "event_count": a.trace.event_count,
        "streams": len(a.trace.streams),
        "prefill_segments": len(a.meta["prefill_segments"]),
        "pd_segments": len(a.meta["pd_segments"]),
        "decode_samples": a.meta["decode_num"],
        "decode_excluded_segments": len(a.meta["decode_excluded"]),
        "stages_ms": {s: round(a.stage_us[s] / 1000, 3) for s in ordered_stages(a)},
        "top_classes": [{"class": c, "ms": round(v / 1000, 3), "pct": round(_pct(v, a.gpu_us), 2)}
                        for c, v in top_items(a.class_us, top)],
        "top_kernels": [{"kernel": f"{c}/{t}/{n}",
                         "ms": round(v / 1000, 3),
                         "pct": round(_pct(v, a.gpu_us), 2),
                         "count": a.kernel_cnt[(c, t, n)]}
                        for (c, t, n), v in top_items(a.kernel_us, top)],
        "warnings": sorted(set(a.meta["warnings"])),
    }


def write_summary_json(a: Analysis, out: Path, top: int = 5) -> None:
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(json.dumps(_summary_payload(a, top), ensure_ascii=False, indent=2),
                   encoding="utf-8")


def write_batch_summary_json(analyses: List[Analysis], out: Path, top: int = 5) -> None:
    payload = {
        "case_count": len(analyses),
        "decode_samples": analyses[0].meta["decode_num"] if analyses else 0,
        "cases": [_summary_payload(a, top) for a in analyses],
    }
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")


# ---------------------------------------------------------------------------
# 多 case 合并输出（一份 md + 一份 xlsx，每张表带 case 列）
# ---------------------------------------------------------------------------
def _common_stem(inputs: List[str]) -> str:
    """多 case 产物名：所有输入同父目录时用父目录名，否则 multi_case"""
    parents = {os.path.dirname(os.path.abspath(i.strip().rstrip("/"))) for i in inputs}
    if len(parents) == 1:
        base = os.path.basename(parents.pop())
        if base and base not in (".", os.sep):
            return base
    return "multi_case"


def _batch_note(a: Analysis, stage: str) -> str:
    n = a.meta["decode_num"]
    return {
        "prefilling": f"全部纯 prefill（{len(a.meta['prefill_segments'])} 段）",
        "prefilling+decoding": "ctx>0 且 gen>0 的混合段",
        "decoding": f"最后 {n} 个纯 decode（采样）",
        "before-dec": "采样窗口内非阶段 kernel",
        "before-pre": "首个阶段之前的 kernel",
        "after-dec": "最后采样阶段之后的 kernel",
        "decoding(excluded)": "更早 decode 段（不计入阶段口径）",
    }.get(stage, "")


def _b_summary(wb, analyses: List[Analysis]) -> None:
    ws = wb.create_sheet("Summary")
    ws.append(["case", "source", "gpu_ms", "span_ms", "busy_pct", "kernel_count",
               "event_count", "streams", "prefill_segments", "pd_segments",
               "decode_samples", "decode_excluded_segments",
               "prefilling_ms", "pd_ms", "decoding_ms"])
    for a in analyses:
        ws.append([
            a.name, a.source, round(a.gpu_us / 1000, 3), round(a.span_us / 1000, 3),
            round(_pct(a.gpu_us, a.span_us), 2), a.kernel_count, a.trace.event_count,
            len(a.trace.streams), len(a.meta["prefill_segments"]),
            len(a.meta["pd_segments"]), a.meta["decode_num"],
            len(a.meta["decode_excluded"]),
            round(a.stage_us.get("prefilling", 0.0) / 1000, 3),
            round(a.stage_us.get("prefilling+decoding", 0.0) / 1000, 3),
            round(a.stage_us.get("decoding", 0.0) / 1000, 3),
        ])
    _style_header(ws, [30, 46, 12, 12, 10, 12, 12, 8, 16, 12, 14, 22, 14, 12, 14])


def _b_stages(wb, analyses: List[Analysis]) -> None:
    ws = wb.create_sheet("Stages")
    ws.append(["case", "stage", "kernel_count", "kernel_us", "kernel_ms", "span_us",
               "pct_of_gpu", "note"])
    for a in analyses:
        for s in ordered_stages(a):
            us = a.stage_us[s]
            ws.append([a.name, s, a.stage_cnt[s], round(us, 3), round(us / 1000, 3),
                       round(a.stage_span.get(s, 0.0), 3) if s in CONTIGUOUS_STAGES else "",
                       round(_pct(us, a.gpu_us), 2), _batch_note(a, s)])
        ws.append([a.name, "TOTAL", sum(a.stage_cnt.values()), round(a.gpu_us, 3),
                   round(a.gpu_us / 1000, 3), round(a.span_us, 3), 100.0,
                   "各阶段之和 == GPU 耗时合计"])
        for c in ws[ws.max_row]:
            c.fill = TOTAL_FILL
    _style_header(ws, [30, 22, 14, 14, 14, 14, 12, 60])


def _b_classes(wb, analyses: List[Analysis]) -> None:
    ws = wb.create_sheet("Classes")
    ws.append(["case", "kernel_class", "kernel_count", "time_us", "time_ms", "pct_of_gpu",
               "prefill_us", "prefill_ms", "pct_of_prefill",
               "decode_us", "decode_ms", "pct_of_decode",
               "pd_us", "pd_ms", "main_stage"])
    for a in analyses:
        pre_total = a.stage_us.get("prefilling", 0.0)
        dec_total = a.stage_us.get("decoding", 0.0)
        for cls, us in sorted(a.class_us.items(), key=lambda kv: -kv[1]):
            pre, dec, pd = _class_split_us(a, cls)
            by_stage = {s: a.stage_class_us.get((s, cls), 0.0) for s in a.stage_us}
            main_stage = max(by_stage.items(), key=lambda kv: kv[1])[0] if by_stage else ""
            ws.append([a.name, cls, a.class_cnt[cls], round(us, 3), round(us / 1000, 3),
                       round(_pct(us, a.gpu_us), 2),
                       round(pre, 3), round(pre / 1000, 3), round(_pct(pre, pre_total), 2),
                       round(dec, 3), round(dec / 1000, 3), round(_pct(dec, dec_total), 2),
                       round(pd, 3), round(pd / 1000, 3), main_stage])
        pre_all = a.stage_us.get("prefilling", 0.0)
        dec_all = a.stage_us.get("decoding", 0.0)
        pd_all = a.stage_us.get("prefilling+decoding", 0.0)
        ws.append([a.name, "TOTAL", sum(a.class_cnt.values()), round(a.gpu_us, 3),
                   round(a.gpu_us / 1000, 3), 100.0,
                   round(pre_all, 3), round(pre_all / 1000, 3), 100.0,
                   round(dec_all, 3), round(dec_all / 1000, 3), 100.0,
                   round(pd_all, 3), round(pd_all / 1000, 3), ""])
        for c in ws[ws.max_row]:
            c.fill = TOTAL_FILL
    _style_header(ws, [30, 26, 14, 14, 14, 12, 14, 14, 14, 14, 14, 14, 14, 14, 22])


def _b_types(wb, analyses: List[Analysis]) -> None:
    ws = wb.create_sheet("Types")
    ws.append(["case", "kernel_class", "kernel_type", "kernel_count", "time_us",
               "time_ms", "pct_of_class", "pct_of_gpu", "main_stage"])
    for a in analyses:
        for (cls, typ), us in sorted(a.type_us.items(), key=lambda kv: (kv[0][0], -kv[1])):
            by_stage = {s: a.stage_type_us.get((s, cls, typ), 0.0) for s in a.stage_us}
            main_stage = max(by_stage.items(), key=lambda kv: kv[1])[0] if by_stage else ""
            ws.append([a.name, cls, typ, a.type_cnt[(cls, typ)], round(us, 3),
                       round(us / 1000, 3), round(_pct(us, a.class_us[cls]), 2),
                       round(_pct(us, a.gpu_us), 2), main_stage])
    _style_header(ws, [30, 22, 22, 14, 14, 14, 14, 12, 22])


def _b_kernels(wb, analyses: List[Analysis]) -> None:
    ws = wb.create_sheet("Kernels")
    ws.append(["case", "kernel_class", "kernel_type", "kernel_name", "kernel_count",
               "time_us", "time_ms", "pct_of_gpu", "main_stage"])
    for a in analyses:
        per_stage: Dict[Tuple[str, str, str], Counter] = defaultdict(Counter)
        for (stage, cls, typ, nm), us in a.stage_kernel_us.items():
            per_stage[(cls, typ, nm)][stage] += us
        for (cls, typ, nm), us in sorted(a.kernel_us.items(), key=lambda kv: -kv[1]):
            stages = per_stage.get((cls, typ, nm), Counter())
            main_stage = stages.most_common(1)[0][0] if stages else ""
            ws.append([a.name, cls, typ, nm, a.kernel_cnt[(cls, typ, nm)], round(us, 3),
                       round(us / 1000, 3), round(_pct(us, a.gpu_us), 2), main_stage])
    _style_header(ws, [30, 22, 22, 70, 14, 14, 14, 12, 22])


def _b_hierarchy(wb, analyses: List[Analysis]) -> None:
    ws = wb.create_sheet("Hierarchy")
    ws.append(["case", "level", "stage", "kernel_class", "kernel_type", "kernel_name",
               "kernel_count", "time_us", "time_ms", "pct_of_stage", "pct_of_class",
               "pct_of_type", "pct_of_gpu"])

    for a in analyses:
        gpu_us = a.gpu_us

        def emit(level, stage, cls, typ, nm, cnt, us, p_stage, p_class=None, p_type=None):
            ws.append([a.name, level, stage, cls, typ, nm, cnt, round(us, 3),
                       round(us / 1000, 3),
                       round(p_stage, 4) if p_stage is not None else "",
                       round(p_class, 4) if p_class is not None else "",
                       round(p_type, 4) if p_type is not None else "",
                       round(_pct(us, gpu_us), 4)])
            if level == 0:
                for c in ws[ws.max_row]:
                    c.fill = GRP_FILL
                    c.font = Font(bold=True)
            elif nm == "total":
                for c in ws[ws.max_row]:
                    c.fill = TOTAL_FILL

        for stage in ordered_stages(a):
            s_us = a.stage_us[stage]
            emit(0, stage, "total", "total", "total", a.stage_cnt[stage], s_us, 100.0)
            classes = sorted(((c, v) for (s, c), v in a.stage_class_us.items() if s == stage),
                             key=lambda kv: -kv[1])
            for cls, us in classes:
                emit(1, stage, cls, "total", "total",
                     a.stage_class_cnt[(stage, cls)], us, _pct(us, s_us), 100.0)
                types = sorted(((k, v) for k, v in a.type_us.items() if k[0] == cls),
                               key=lambda kv: -kv[1])
                for (c2, t), t_us in types:
                    st_us = a.stage_type_us.get((stage, c2, t), 0.0)
                    if st_us <= 0:
                        continue
                    emit(2, stage, cls, t, "total",
                         a.stage_type_cnt.get((stage, c2, t), 0), st_us,
                         _pct(st_us, s_us), _pct(st_us, us), 100.0)
                    kernels = sorted(((k, v) for k, v in a.kernel_us.items()
                                      if k[0] == cls and k[1] == t),
                                     key=lambda kv: -kv[1])
                    for (c3, t3, nm), _ in kernels:
                        k_stage_us = a.stage_kernel_us.get((stage, c3, t3, nm), 0.0)
                        if k_stage_us <= 0:
                            continue
                        emit(3, stage, cls, t, nm,
                             a.stage_kernel_cnt.get((stage, c3, t3, nm), 0), k_stage_us,
                             _pct(k_stage_us, s_us), _pct(k_stage_us, us),
                             _pct(k_stage_us, t_us))
    _style_header(ws, [30, 8, 22, 20, 20, 70, 14, 14, 14, 12, 12, 12, 12])


def write_batch_excel(analyses: List[Analysis], out: Path) -> None:
    wb = Workbook()
    wb.remove(wb.active)
    _b_summary(wb, analyses)
    _b_stages(wb, analyses)
    _b_classes(wb, analyses)
    _b_types(wb, analyses)
    _b_kernels(wb, analyses)
    _b_hierarchy(wb, analyses)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)


def write_batch_markdown(analyses: List[Analysis], out: Path, top: int) -> None:
    n_case = len(analyses)
    n_dec = analyses[0].meta["decode_num"] if analyses else 0

    L: List[str] = []
    L.append(f"# Torch Profile 分析报告（{n_case} 个 Case）")
    L.append("")
    L.append(f"> 生成时间：{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}  ")
    L.append(f"> 分析工具：`tpa.py`（ModelzooTool `trace-analyzer` skill）  ")
    L.append(f"> decode 采样：每个 case 各取**最后 {n_dec} 个纯 decode 段**")
    L.append("")
    L.append("## 0. 总览")
    L.append("")
    L.append("| Case | GPU 耗时 (ms) | 覆盖时长 (ms) | 忙碌率 | kernel 数 | prefill | decode（采样 "
             f"{n_dec} 段） | prefill+decode |")
    L.append("| --- | ---: | ---: | ---: | ---: | ---: | ---: | ---: |")
    for a in analyses:
        pd_us = a.stage_us.get("prefilling+decoding", 0.0)
        L.append(f"| `{a.name}` | **{_fmt_ms(a.gpu_us)}** | {_fmt_ms(a.span_us)} | "
                 f"{_pct(a.gpu_us, a.span_us):.2f}% | {a.kernel_count:,} | "
                 f"{len(a.meta['prefill_segments'])} 段 / "
                 f"{_fmt_ms(a.stage_us.get('prefilling', 0.0))} ms | "
                 f"{_fmt_ms(a.stage_us.get('decoding', 0.0))} ms | "
                 f"{(_fmt_ms(pd_us) + ' ms') if pd_us else '-'} |")
    L.append("")

    for idx, a in enumerate(analyses, start=1):
        n = a.meta["decode_num"]
        dec_segs = a.meta["decode_segments"]
        dec_step = (sum(s.end - s.start for s in dec_segs) / n / 1000) if n else 0.0
        L.append(f"## {idx}. Case：`{a.name}`")
        L.append("")
        L.append(f"- 数据来源：`{a.source}`")
        L.append(f"- GPU 耗时 **{_fmt_ms(a.gpu_us)} ms** ｜ 覆盖时长 {_fmt_ms(a.span_us)} ms ｜ "
                 f"忙碌率 {_pct(a.gpu_us, a.span_us):.2f}% ｜ kernel {a.kernel_count:,} 个")
        L.append(f"- prefill {len(a.meta['prefill_segments'])} 段；"
                 f"decode 采样 {n} 段，合计 {_fmt_ms(a.stage_us.get('decoding', 0.0))} ms，"
                 f"平均 {_fmt_ms(a.stage_us.get('decoding', 0.0) / n if n else 0.0)} ms/段"
                 f"（单步 wall-clock ≈ {dec_step:.2f} ms）")
        if a.meta["decode_excluded"]:
            L.append(f"- 另有 {len(a.meta['decode_excluded'])} 个更早 decode 段未纳入阶段统计"
                     f"（仅计入全量）")
        L.append("")

        L.append("**阶段耗时**")
        L.append("")
        L.append("| stage | kernel 数 | 耗时 (ms) | 占 GPU | 说明 |")
        L.append("| --- | ---: | ---: | ---: | --- |")
        for s in ordered_stages(a):
            L.append(f"| {s} | {a.stage_cnt[s]:,} | {_fmt_ms(a.stage_us[s])} | "
                     f"{_pct(a.stage_us[s], a.gpu_us):.2f}% | {_batch_note(a, s)} |")
        L.append(f"| **合计** | **{a.kernel_count:,}** | **{_fmt_ms(a.gpu_us)}** | "
                 f"**100%** | 之和 == GPU 耗时合计 |")
        L.append("")

        L.append(f"**算子大类（Top {top}，prefill / decode 拆分）**")
        L.append("")
        L.append("| 大类 | 总耗时 (ms) | 占 GPU | prefill (ms) | 占 prefill | "
                 "decode (ms) | 占 decode |")
        L.append("| --- | ---: | ---: | ---: | ---: | ---: | ---: |")
        pre_total = a.stage_us.get("prefilling", 0.0)
        dec_total = a.stage_us.get("decoding", 0.0)
        for cls, us in top_items(a.class_us, top):
            pre, dec, _ = _class_split_us(a, cls)
            L.append(f"| {cls} | {_fmt_ms(us)} | {_pct(us, a.gpu_us):.2f}% | "
                     f"{_fmt_ms(pre)} | {_pct(pre, pre_total):.2f}% | "
                     f"{_fmt_ms(dec)} | {_pct(dec, dec_total):.2f}% |")
        pd_us = a.stage_us.get("prefilling+decoding", 0.0)
        if pd_us:
            L.append(f"\n> 本 case 另有 **{_fmt_ms(pd_us)} ms**（占 GPU "
                     f"{_pct(pd_us, a.gpu_us):.2f}%）落在 `prefilling+decoding` 混合段，"
                     f"未计入上述 prefill / decode 两列。")
        L.append("")

        L.append(f"**算子子类（Top {top}）**")
        L.append("")
        L.append("| 大类 | 子类 | 耗时 (ms) | 占本类 | 占 GPU |")
        L.append("| --- | --- | ---: | ---: | ---: |")
        for (cls, typ), us in top_items(a.type_us, top):
            L.append(f"| {cls} | {typ} | {_fmt_ms(us)} | {_pct(us, a.class_us[cls]):.2f}% | "
                     f"{_pct(us, a.gpu_us):.2f}% |")
        L.append("")

        L.append(f"**单算子（Top {top}）**")
        L.append("")
        L.append("| 大类/子类 | 算子 | 耗时 (ms) | 占 GPU | 次数 |")
        L.append("| --- | --- | ---: | ---: | ---: |")
        for (cls, typ, nm), us in top_items(a.kernel_us, top):
            short = nm if len(nm) <= 56 else nm[:53] + "..."
            L.append(f"| {cls}/{typ} | `{short}` | {_fmt_ms(us)} | "
                     f"{_pct(us, a.gpu_us):.2f}% | {a.kernel_cnt[(cls, typ, nm)]:,} |")
        L.append("")

    L.append("## 附：统计口径与校验")
    L.append("")
    L.append("1. kernel 来源：trace 中 `cat == \"kernel\"` 的事件，耗时字段为 `dur`（微秒）。")
    L.append("2. 阶段来源：`cat == \"gpu_user_annotation\"` 且 name 以 `execute_` 开头的注解，"
             "按 `context_N(...)` / `generation_M(...)` 判定：`N>0 且 M==0` → prefill，"
             "`N==0 且 M>0` → decode，两者皆 >0 → 混合段。")
    L.append(f"3. **decode 采样数量为 {n_dec}**（每个 case 独立取最后 {n_dec} 个纯 decode 段）；"
             "更早的 decode 段标记 `decoding(excluded)`，只计入全量。")
    L.append("4. prefill 统计全部纯 prefill 段（chunked prefill 会拆成多段）。")
    L.append("5. 落在注解时间窗之外的 kernel 归入 `before-pre` / `before-dec` / `after-dec`，不会丢失。")
    L.append("6. 校验（各 stage 之和 == GPU 耗时合计）：")
    L.append("")
    L.append("| Case | 各阶段之和 (ms) | GPU 耗时合计 (ms) | 一致 |")
    L.append("| --- | ---: | ---: | --- |")
    for a in analyses:
        ssum = sum(a.stage_us.values())
        L.append(f"| `{a.name}` | {_fmt_ms(ssum)} | {_fmt_ms(a.gpu_us)} | "
                 f"{'✅' if abs(ssum - a.gpu_us) < 1e-6 else '❌'} |")
    L.append("")

    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text("\n".join(L), encoding="utf-8")


# ---------------------------------------------------------------------------
# 对比模式（保留原项目能力）
# ---------------------------------------------------------------------------
def _hier_rows(a: Analysis) -> Dict[Tuple[str, str, str, str], Tuple[int, float]]:
    """(stage, class, type, kernel) -> (count, us)，'total' 表示上层合计"""
    rows: Dict[Tuple[str, str, str, str], Tuple[int, float]] = {}
    for (stage, cls, typ, nm), us in a.stage_kernel_us.items():
        rows[(stage, cls, typ, nm)] = (a.stage_kernel_cnt[(stage, cls, typ, nm)], us)
    for (stage, cls, typ), us in a.stage_type_us.items():
        rows[(stage, cls, typ, "total")] = (a.stage_type_cnt[(stage, cls, typ)], us)
    for (stage, cls), us in a.stage_class_us.items():
        rows[(stage, cls, "total", "total")] = (a.stage_class_cnt[(stage, cls)], us)
    for stage in ordered_stages(a):
        rows[(stage, "total", "total", "total")] = (a.stage_cnt[stage], a.stage_us[stage])
    return rows


def write_compare(ctl: Analysis, exp: Analysis, out_xlsx: Path, out_md: Path, top: int) -> None:
    ctl_rows, exp_rows = _hier_rows(ctl), _hier_rows(exp)
    keys = sorted(set(ctl_rows) | set(exp_rows),
                  key=lambda k: (STAGE_ORDER.index(k[0]) if k[0] in STAGE_ORDER else 99,
                                 k[0], k[1], k[2], k[3]))

    wb = Workbook()
    ws = wb.active
    ws.title = "Compare"
    ws.append(["Case_Name", "stage", "kernel_class", "kernel_type", "EXP_Kernel_Name",
               "EXP_Count", "EXP_Time", "CTL_Kernel_Name", "CTL_Count", "CTL_Time",
               "Performance[kernel_class|kernel_type](%)", "EXP_Time_ms", "CTL_Time_ms", "Delta_ms"])
    for k in keys:
        stage, cls, typ, nm = k
        ec, eu = exp_rows.get(k, (0, 0.0))
        cc, cu = ctl_rows.get(k, (0, 0.0))
        perf = (cu / eu * 100) if eu > 0 else None
        ws.append([
            f"{ctl.name}__vs__{exp.name}", stage, cls, typ, nm,
            ec, round(eu, 3), nm, cc, round(cu, 3),
            (f"{perf:.2f}%" if perf is not None else "NaN"),
            round(eu / 1000, 3), round(cu / 1000, 3), round((cu - eu) / 1000, 3),
        ])
        if nm == "total":
            for c in ws[ws.max_row]:
                c.fill = TOTAL_FILL
        elif typ == "total":
            for c in ws[ws.max_row]:
                c.fill = GRP_FILL
    _style_header(ws, [30, 22, 20, 20, 60, 12, 14, 60, 12, 14, 30, 14, 14, 12])

    ws2 = wb.create_sheet("Summary")
    ws2.append(["指标", f"CTL({ctl.name})", f"EXP({exp.name})", "EXP/CTL"])
    pairs = [
        ("GPU 耗时合计 (ms)", ctl.gpu_us / 1000, exp.gpu_us / 1000),
        ("profile 覆盖时长 (ms)", ctl.span_us / 1000, exp.span_us / 1000),
        ("kernel 事件数", ctl.kernel_count, exp.kernel_count),
    ]
    for s in ordered_stages(ctl) + [s for s in ordered_stages(exp) if s not in ctl.stage_us]:
        pairs.append((f"{s} (ms)", ctl.stage_us.get(s, 0.0) / 1000, exp.stage_us.get(s, 0.0) / 1000))
    for name, c, e in pairs:
        ratio = (e / c) if c else None
        ws2.append([name, round(c, 3), round(e, 3), round(ratio, 4) if ratio else ""])
    _style_header(ws2, [30, 18, 18, 12])

    ws3 = wb.create_sheet("Classes")
    ws3.append(["kernel_class", "CTL_ms", "EXP_ms", "Delta_ms", "Performance(CTL/EXP %)"])
    for cls in sorted(set(ctl.class_us) | set(exp.class_us),
                      key=lambda c: -(ctl.class_us.get(c, 0) + exp.class_us.get(c, 0))):
        cu, eu = ctl.class_us.get(cls, 0.0), exp.class_us.get(cls, 0.0)
        ws3.append([cls, round(cu / 1000, 3), round(eu / 1000, 3), round((cu - eu) / 1000, 3),
                    round(cu / eu * 100, 2) if eu else "NaN"])
    _style_header(ws3, [24, 16, 16, 14, 24])

    out_xlsx.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_xlsx)

    # Markdown
    L: List[str] = []
    L.append(f"# Torch Profile 对比报告：{ctl.name} (CTL) vs {exp.name} (EXP)")
    L.append("")
    L.append("> `Performance = CTL_Time / EXP_Time × 100%`；> 100% 表示 EXP 更快（耗时更少），< 100% 表示变慢。")
    L.append(f"> decode 采样数量：CTL {ctl.meta['decode_num']} 段 / EXP {exp.meta['decode_num']} 段。")
    L.append("")
    L.append("## 1. 概览")
    L.append("")
    L.append("| 指标 | CTL | EXP | CTL/EXP |")
    L.append("| --- | ---: | ---: | ---: |")
    L.append(f"| GPU 耗时合计 (ms) | {_fmt_ms(ctl.gpu_us)} | {_fmt_ms(exp.gpu_us)} | "
             f"{ctl.gpu_us/exp.gpu_us:.3f} |" if exp.gpu_us else
             f"| GPU 耗时合计 (ms) | {_fmt_ms(ctl.gpu_us)} | {_fmt_ms(exp.gpu_us)} | - |")
    L.append(f"| profile 覆盖时长 (ms) | {_fmt_ms(ctl.span_us)} | {_fmt_ms(exp.span_us)} | - |")
    L.append(f"| kernel 事件数 | {ctl.kernel_count:,} | {exp.kernel_count:,} | - |")
    L.append("")
    L.append("## 2. 阶段耗时对比")
    L.append("")
    L.append("| stage | CTL (ms) | EXP (ms) | Delta (ms) | 变化 |")
    L.append("| --- | ---: | ---: | ---: | ---: |")
    for s in STAGE_ORDER:
        if s not in ctl.stage_us and s not in exp.stage_us:
            continue
        cu, eu = ctl.stage_us.get(s, 0.0), exp.stage_us.get(s, 0.0)
        delta = eu - cu
        chg = f"{(delta / cu * 100):+.2f}%" if cu else "-"
        L.append(f"| {s} | {_fmt_ms(cu)} | {_fmt_ms(eu)} | {_fmt_ms(delta)} | {chg} |")
    L.append("")

    L.append(f"## 3. 算子大类对比（按 EXP 耗时排序，Top {top}）")
    L.append("")
    L.append("| 大类 | CTL (ms) | EXP (ms) | Delta (ms) | Performance |")
    L.append("| --- | ---: | ---: | ---: | ---: |")
    allcls = sorted(set(ctl.class_us) | set(exp.class_us), key=lambda c: -exp.class_us.get(c, 0))
    for cls in allcls[:top]:
        cu, eu = ctl.class_us.get(cls, 0.0), exp.class_us.get(cls, 0.0)
        perf = f"{cu/eu*100:.2f}%" if eu else "NaN"
        L.append(f"| {cls} | {_fmt_ms(cu)} | {_fmt_ms(eu)} | {_fmt_ms(eu-cu)} | {perf} |")
    L.append("")

    L.append(f"## 4. 单算子差异（|Delta| Top {top}）")
    L.append("")
    L.append("| 大类/子类 | 算子 | CTL (ms) | EXP (ms) | Delta (ms) | Performance |")
    L.append("| --- | --- | ---: | ---: | ---: | ---: |")
    diffs = []
    for (cls, typ, nm), eu in exp.kernel_us.items():
        cu = ctl.kernel_us.get((cls, typ, nm), 0.0)
        diffs.append((abs(eu - cu), cls, typ, nm, cu, eu))
    for _, cls, typ, nm, cu, eu in sorted(diffs, key=lambda x: -x[0])[:top]:
        perf = f"{cu/eu*100:.2f}%" if eu else "NaN"
        short = nm if len(nm) <= 56 else nm[:53] + "..."
        L.append(f"| {cls}/{typ} | `{short}` | {_fmt_ms(cu)} | {_fmt_ms(eu)} | "
                 f"{_fmt_ms(eu-cu)} | {perf} |")
    L.append("")
    L.append("## 5. 口径说明")
    L.append("")
    L.append(f"1. 两侧各自独立解析，按 `(stage, kernel_class, kernel_type, kernel_name)` 对齐；"
             f"未匹配到的一侧记 0。")
    L.append(f"2. decode 采样：CTL 取最后 {ctl.meta['decode_num']} 个纯 decode 段，"
             f"EXP 取最后 {exp.meta['decode_num']} 个纯 decode 段。")
    L.append("3. 耗时单位为微秒累加（Excel `*_Time` 列沿用原项目 us 口径），Markdown 中统一换算为 ms。")
    L.append("")

    out_md.parent.mkdir(parents=True, exist_ok=True)
    out_md.write_text("\n".join(L), encoding="utf-8")


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------
def _default_name_from_source(src: str) -> str:
    """从输入推导产物名：
    - 压缩包成员（dir/x.zip!member）→ 优先用容器的父目录名
    - 形如 dp0_pp0_tp0_..._rank0 的文件名 → 回退到父目录名（保留业务信息）
    """
    suffixes = r"\.(tar\.gz|tar\.bz2|tar\.xz|tgz|tbz2|txz|tar|zip|gz|json)$"

    def strip(name: str) -> str:
        # 循环剥离：.pt.trace.json.gz → .pt.trace.json → .pt.trace → ""
        while True:
            new = re.sub(suffixes, "", name, flags=re.IGNORECASE)
            if new == name:
                break
            name = new
        name = re.sub(r"\.pt\.trace$", "", name, flags=re.IGNORECASE)
        name = re.sub(r"\.\d{15,}$", "", name)              # trace 时间戳后缀
        return name

    container, _, member = src.partition("!")
    target = (member or container).rstrip("/")
    base = strip(os.path.basename(target))

    rank_pat = re.compile(r"^(dp\d+_)?(pp\d+_)?(tp\d+_)?(dcp\d+_)?(ep\d+_)?rank\d+$")
    if rank_pat.match(base):
        # 文件名只剩 rank 信息：改用父目录 / 容器名
        parent = os.path.basename(os.path.dirname(os.path.abspath(target)))
        if member:
            base = strip(os.path.basename(container)) or base
        elif parent and not rank_pat.match(parent):
            base = parent
    return base or "trace"


def _resolve_candidate(target: str, rank: Optional[int], trace_substr: Optional[str]) -> Candidate:
    cands = collect_candidates(target)
    cand = select_candidate(cands, rank, trace_substr)
    gpu_cnt = sum(1 for c in cands if "async_llm" not in c.label.lower())
    if gpu_cnt > 1:
        print(f"{YELLOW}检测到 {gpu_cnt} 个 GPU trace（多 rank）。本次使用：{cand.label}"
              f"；如需指定请用 --rank N 或 --trace SUBSTR{NC}")
    print(f"选中: {cand.label}  (rank={cand.rank}, {cand.size/1024/1024:.2f}MB)")
    return cand


def main(argv: Optional[List[str]] = None) -> None:
    ap = argparse.ArgumentParser(
        description="Torch Profile 分析：算子分类统计 + 阶段耗时 + Markdown/Excel 报告（支持压缩包，无需解压）",
        formatter_class=argparse.ArgumentDefaultsHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)

    for name, help_text in (("analyze", "单组统计分析"),
                            ("compare", "对照组 vs 实验组对比"),
                            ("list", "列出输入中的可用 trace")):
        p = sub.add_parser(name, help=help_text)
        if name == "compare":
            p.add_argument("ctl", help="对照组输入（文件/压缩包/目录）")
            p.add_argument("exp", help="实验组输入（文件/压缩包/目录）")
            p.add_argument("--rank-ctl", type=int, default=None)
            p.add_argument("--rank-exp", type=int, default=None)
        elif name == "analyze":
            p.add_argument("inputs", nargs="+",
                           help="一个或多个输入（文件/压缩包/目录）；多个 case 会合并成一份报告")
        else:
            p.add_argument("input", help="输入（trace 文件 / .gz / zip / tar.gz / 目录）")
        if name == "list":
            p.add_argument("--json", action="store_true",
                           help="以 JSON 输出候选列表（供 agent 生成 rank 选项）")
        p.add_argument("--rank", type=int, default=0, help="优先选择的 rank 编号")
        p.add_argument("--trace", default=None, help="按文件名子串选择 trace")
        p.add_argument("--decode-num", type=int, default=5, help="decode 采样段数")
        p.add_argument("--outdir", default="tpa_out", help="输出目录")
        p.add_argument("--name", default=None, help="产物命名前缀")
        p.add_argument("--top", type=int, default=15, help="Markdown 表格行数")
        p.add_argument("--summary-json", default=None, help="额外输出极简摘要 JSON")

    args = ap.parse_args(argv)

    if args.cmd == "list":
        cmd_list(args.input, getattr(args, "json", False))
        return

    outdir = Path(args.outdir).expanduser().resolve()
    outdir.mkdir(parents=True, exist_ok=True)

    if args.cmd == "analyze":
        analyses: List[Analysis] = []
        for target in args.inputs:
            cand = _resolve_candidate(target, args.rank, args.trace)
            analyses.append(analyze(cand, _default_name_from_source(target.strip()),
                                    args.decode_num))

        if len(analyses) == 1:
            a = analyses[0]
            md = outdir / f"{a.name}_report.md"
            xlsx = outdir / f"{a.name}.xlsx"
            write_markdown(a, md, args.top)
            write_excel(a, xlsx)
            print(f"\n{GREEN}报告: {md}\nExcel: {xlsx}{NC}")
            if args.summary_json:
                write_summary_json(a, Path(args.summary_json))
                print(f"摘要: {args.summary_json}")
            print(f"GPU 耗时合计 {a.gpu_us/1000:,.2f} ms，"
                  f"prefill {len(a.meta['prefill_segments'])} 段 / "
                  f"decode 采样 {a.meta['decode_num']} 段")
            return

        # 多 case：合并成一份 md + 一份 xlsx，避免产物碎散
        stem = args.name or _common_stem(args.inputs)
        md = outdir / f"{stem}_report.md"
        xlsx = outdir / f"{stem}.xlsx"
        write_batch_markdown(analyses, md, args.top)
        write_batch_excel(analyses, xlsx)
        print(f"\n{GREEN}合并报告（{len(analyses)} 个 case）: {md}\n合并 Excel: {xlsx}{NC}")
        if args.summary_json:
            write_batch_summary_json(analyses, Path(args.summary_json), args.top)
            print(f"摘要: {args.summary_json}")
        for a in analyses:
            print(f"  - {a.name}: GPU {a.gpu_us/1000:,.2f} ms，"
                  f"prefill {len(a.meta['prefill_segments'])} 段 / "
                  f"decode 采样 {a.meta['decode_num']} 段")
        return

    # compare
    ctl_cand = _resolve_candidate(args.ctl, args.rank_ctl if args.rank_ctl is not None else args.rank,
                                  args.trace)
    exp_cand = _resolve_candidate(args.exp, args.rank_exp if args.rank_exp is not None else args.rank,
                                  args.trace)
    ctl_name = _default_name_from_source(args.ctl.strip())
    exp_name = _default_name_from_source(args.exp.strip())
    ctl = analyze(ctl_cand, ctl_name, args.decode_num)
    exp = analyze(exp_cand, exp_name, args.decode_num)

    stem = args.name or f"{ctl_name}__vs__{exp_name}"
    md = outdir / f"{stem}_compare.md"
    xlsx = outdir / f"{stem}_compare.xlsx"
    write_compare(ctl, exp, xlsx, md, args.top)
    print(f"\n{GREEN}对比报告: {md}\n对比 Excel: {xlsx}{NC}")
    if args.summary_json:
        payload = {
            "ctl": {"name": ctl.name, "gpu_ms": round(ctl.gpu_us / 1000, 3)},
            "exp": {"name": exp.name, "gpu_ms": round(exp.gpu_us / 1000, 3)},
            "performance_pct": round(ctl.gpu_us / exp.gpu_us * 100, 2) if exp.gpu_us else None,
        }
        Path(args.summary_json).parent.mkdir(parents=True, exist_ok=True)
        Path(args.summary_json).write_text(json.dumps(payload, ensure_ascii=False, indent=2),
                                           encoding="utf-8")
        print(f"摘要: {args.summary_json}")


if __name__ == "__main__":
    main()
